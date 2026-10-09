"""Cảnh báo, KPI, bản tin chống bịa số, Text-to-SQL an toàn, Excel và email."""

import smtplib

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from src import anomaly, ask_data, briefs, deliver, departments, excel_report, finance, kpis, metrics
from src.pipeline import recovery_list


# ----------------------------------------------------------------- cảnh báo bất thường
def weekly_frame(external_rates, n=1000):
    weeks = pd.date_range("2022-01-03", periods=len(external_rates), freq="W-MON")
    df = pd.DataFrame({"n_tickets": n, "external": np.array(external_rates) * n, "customer": 0.02 * n,
                       "internal": 0.001 * n}, index=weeks)
    for g in anomaly.ERROR_GROUPS:
        df[f"{g}_rate"] = df[g] / df["n_tickets"]
    return df


def test_spike_is_flagged_and_drop_is_not():
    rates = [0.08, 0.081, 0.079, 0.08, 0.082, 0.078, 0.08, 0.081, 0.18, 0.03]
    alerts = anomaly.detect(weekly_frame(rates))
    ext = alerts[alerts["error_group"] == "external"].set_index("week")["is_alert"]
    assert ext.iloc[8] and not ext.iloc[9] and ext.sum() == 1


def test_small_weeks_are_ignored():
    alerts = anomaly.detect(weekly_frame([0.08] * 8 + [0.5], n=50))
    assert alerts.empty


def test_real_2022_bank_incident_is_flagged(tickets):
    alerts = anomaly.detect(metrics.weekly_error_rates(tickets))
    hit = alerts[(alerts["week"] == "2022-02-28") & (alerts["error_group"] == "external")].iloc[0]
    assert hit["is_alert"] and hit["z"] > 4


# ----------------------------------------------------------------- KPI / tài chính
def test_kpi_rates_are_consistent(tickets):
    k = kpis.monthly_kpis(tickets)
    assert np.allclose(k["payment_success_rate"] + k["payment_failure_rate"], 1)
    assert k.loc["2022-12", "repeat_90d"] != k.loc["2022-12", "repeat_90d"]  # NaN: chưa đủ 90 ngày quan sát


def test_finance_hides_incomplete_months(tickets):
    losses = finance.monthly_losses(tickets)
    assert pd.isna(losses.loc["2022-10", "discount_one_time"])       # cần 90 ngày sau tháng 10
    assert pd.notna(losses.loc["2022-09", "discount_one_time"])


# ----------------------------------------------------------------- bản tin chống bịa số
@pytest.fixture(scope="module")
def week_facts(tickets):
    alerts = anomaly.detect(metrics.weekly_error_rates(tickets))
    facts = briefs.weekly_facts(tickets, alerts, "2022-02-28")
    extra, tables = departments.build(tickets, "2022-02-28")
    facts.update(extra)
    return facts, tables, alerts


def test_validator_catches_invented_numbers(week_facts):
    facts = week_facts[0]
    brief = briefs.template_brief(briefs.facts_for(facts, "finance"), "finance")
    fake = dict(brief, summary=brief["summary"] + " Mất 25,000 doanh thu, lỗi tăng 9.1%.")
    assert briefs.validate_numbers(fake, facts) == ["25,000", "9.1%"]


@pytest.mark.parametrize("dept", list(briefs.DEPARTMENTS))
def test_templates_only_use_fact_numbers(week_facts, dept):
    scoped = briefs.facts_for(week_facts[0], dept)
    assert briefs.validate_numbers(briefs.template_brief(scoped, dept), scoped) == []


def test_brief_falls_back_to_template_without_llm(week_facts):
    _, source = briefs.generate_brief(week_facts[0], "product_it")
    assert source.startswith("template")


# ----------------------------------------------------------------- Text-to-SQL an toàn
@pytest.mark.parametrize("sql", ["DELETE FROM tickets", "SELECT 1; DROP TABLE tickets",
                                 "PRAGMA table_info(tickets)", "UPDATE tickets SET year = 1"])
def test_unsafe_sql_is_blocked(sql):
    with pytest.raises(ValueError):
        ask_data.check_sql(sql)


def test_sql_in_code_fence_is_unwrapped_but_still_checked():
    assert ask_data.check_sql("```sql\nSELECT COUNT(*) FROM tickets;\n```") == "SELECT COUNT(*) FROM tickets"
    with pytest.raises(ValueError):
        ask_data.check_sql("```sql\nDELETE FROM tickets\n```")


def test_sql_grading_allows_extra_columns():
    gold = ask_data._normalize(pd.DataFrame({"movie": ["Avatar"]}))
    assert ask_data._matches(gold, ask_data._normalize(pd.DataFrame({"movie": ["Avatar"], "n": [10]})))
    assert not ask_data._matches(gold, ask_data._normalize(pd.DataFrame({"movie": ["Batman"]})))


# ----------------------------------------------------------------- Excel + email
def test_department_workbook_has_only_its_sheet(week_facts, tickets, tmp_path):
    facts, tables, alerts = week_facts
    generated = {d: briefs.generate_brief(facts, d) for d in briefs.DEPARTMENTS}
    recovery = recovery_list(tickets, pd.Timestamp("2022-02-28"))
    full = excel_report.build(tmp_path / "full.xlsx", facts, generated, tables, recovery, alerts)
    part = excel_report.build(tmp_path / "cs.xlsx", facts, generated, tables, recovery, alerts, departments=["customer_care"])
    assert len(load_workbook(full).sheetnames) == 7
    assert load_workbook(part).sheetnames == ["Tóm tắt", "CSKH", "Chất lượng dữ liệu"]


def test_email_without_smtp_only_saves_previews(week_facts, tmp_path, monkeypatch):
    def never(*_, **__):
        raise AssertionError("không được kết nối SMTP khi chưa cấu hình")
    monkeypatch.setattr(smtplib, "SMTP", never)
    facts = week_facts[0]
    generated = {d: (briefs.template_brief(briefs.facts_for(facts, d), d), "template") for d in briefs.DEPARTMENTS}
    attachment = tmp_path / "a.xlsx"
    attachment.write_bytes(b"x")
    results = deliver.deliver(tmp_path, facts, generated, {d: attachment for d in generated}, recipients={}, settings=None)
    assert {r["status"] for r in results} == {"saved_to_outbox"}
    assert len(list((tmp_path / "outbox").glob("*.eml"))) == 5
