"""Báo cáo vận hành tuần dạng Excel: 1 sheet tóm tắt + 1 sheet / phòng ban + chất lượng dữ liệu.

Người nhận mở bằng Excel là đọc được ngay: số đã định dạng, ô cảnh báo tô đỏ, có biểu đồ Excel gốc
(không phải ảnh) để họ tự lọc / sửa tiếp.
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src import kpis

NAVY = "091540"
FILL_HEADER = PatternFill("solid", fgColor=NAVY)
FILL_RED = PatternFill("solid", fgColor="F8D7DA")
FILL_GREEN = PatternFill("solid", fgColor="D4EDDA")
FILL_GREY = PatternFill("solid", fgColor="F2F4F7")
FONT_HEADER = Font(bold=True, color="FFFFFF")
FONT_TITLE = Font(bold=True, size=16, color=NAVY)
FONT_H2 = Font(bold=True, size=12, color=NAVY)
THIN = Border(bottom=Side(style="thin", color="D0D5DD"))

PCT = "0.0%"
INT = "#,##0"
MONEY = "#,##0.00"

SHEETS = {
    "leadership": "Ban lãnh đạo",
    "marketing": "Marketing",
    "customer_care": "CSKH",
    "product_it": "Product-IT",
    "finance": "Tài chính",
}


def _title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = FONT_TITLE
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="667085")
    ws.sheet_view.showGridLines = False
    return 4


def _h2(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = FONT_H2
    return row + 1


def _brief(ws, row, brief, source):
    row = _h2(ws, row, brief["headline"])
    ws.cell(row=row, column=1, value=f"Nguồn bản tin: {source}").font = Font(italic=True, size=9, color="667085")
    row += 1
    for point in brief["key_points"]:
        cell = ws.cell(row=row, column=1, value=f"• {point}")
        cell.alignment = Alignment(wrap_text=False)
        row += 1
    row += 1
    row = _h2(ws, row, "Hành động đề xuất")
    return _table(ws, row, pd.DataFrame(brief["actions"]).rename(
        columns={"action": "Việc cần làm", "owner": "Phụ trách", "kpi": "Chỉ số theo dõi"}))[1] + 2


def _table(ws, row, df, formats=None, index=False):
    """Ghi DataFrame có header tô màu. Trả về (dòng header, dòng cuối, {tên cột: chữ cột})."""
    formats = formats or {}
    data = df.reset_index() if index else df
    letters = {}
    for j, col in enumerate(data.columns, start=1):
        cell = ws.cell(row=row, column=j, value=str(col))
        cell.fill, cell.font = FILL_HEADER, FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        letters[col] = get_column_letter(j)
    for i, values in enumerate(data.itertuples(index=False), start=row + 1):
        for j, (col, value) in enumerate(zip(data.columns, values), start=1):
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            elif not isinstance(value, (list, dict, str)) and pd.isna(value):
                value = None
            cell = ws.cell(row=i, column=j, value=value)
            cell.border = THIN
            if col in formats:
                cell.number_format = formats[col]
    return row, row + len(data), letters


def _widths(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _summary_sheet(wb, facts, briefs, recovery, dq_status, departments):
    ws = wb.active
    ws.title = "Tóm tắt"
    row = _title(ws, f"Báo cáo vận hành tuần {facts['week_start']} → {facts['week_end']}",
                 "Tạo tự động bởi src/pipeline.py · số liệu chỉ dùng dữ liệu tới hết tuần")
    ext = facts["error_groups"].get("external", {})
    tiles = pd.DataFrame([
        ("Lượt thanh toán", facts["tickets"]),
        ("Tỷ lệ thanh toán thành công", f"{facts['success_rate']} (8 tuần trước: {facts['success_rate_prev_8_weeks']})"),
        ("Giao dịch lỗi", facts["failed_tickets"]),
        ("Khách mới bị mất ở bước thanh toán", facts["new_customers_lost_at_payment"]),
        ("Cảnh báo lỗi phía ngân hàng", f"CÓ - {ext['rate']} (nền {ext['baseline_rate']})" if ext.get("alert") else "Không"),
        ("Danh sách cứu đơn (CSKH)", f"{len(recovery):,} khách"),
        ("Chất lượng dữ liệu", dq_status),
    ], columns=["Chỉ số", "Giá trị"])
    head, last, _ = _table(ws, row, tiles)
    for r in range(head + 1, last + 1):
        value = str(ws.cell(row=r, column=2).value)
        if value.startswith("CÓ") or value.startswith("LỖI"):
            ws.cell(row=r, column=2).fill = FILL_RED
    row = last + 2
    row = _h2(ws, row, "Tiêu điểm từng phòng ban (bấm tên sheet để mở)")
    for j, name in enumerate(["Phòng ban", "Tiêu điểm tuần này"], start=1):
        cell = ws.cell(row=row, column=j, value=name)
        cell.fill, cell.font = FILL_HEADER, FONT_HEADER
    for dept, sheet in SHEETS.items():
        if dept not in departments:
            continue
        row += 1
        link = ws.cell(row=row, column=1, value=sheet)
        link.hyperlink = f"#'{sheet}'!A1"
        link.font = Font(color="0563C1", underline="single")
        ws.cell(row=row, column=2, value=briefs[dept][0]["headline"])
    _widths(ws, {"A": 38, "B": 90})


def _leadership_sheet(wb, facts, brief, tables):
    ws = wb.create_sheet(SHEETS["leadership"])
    row = _title(ws, "Ban lãnh đạo · 6 nhóm KPI", f"KPI tháng {facts['leadership']['kpi_month']} so với tháng trước")
    row = _brief(ws, row, *brief)
    kpi_table = tables["kpis"]
    month = facts["leadership"]["kpi_month"]
    if month in kpi_table.index:
        months = kpi_table.index.tolist()
        prev = months[months.index(month) - 1] if months.index(month) > 0 else None
        rows = []
        for group, key, label, fmt, direction in kpis.KPI_DEFINITIONS:
            now = kpi_table.at[month, key]
            before = kpi_table.at[prev, key] if prev else None
            change = now - before if pd.notna(now) and before is not None and pd.notna(before) else None
            verdict = "" if change is None or change == 0 else ("Tốt hơn" if change * direction > 0 else "Xấu hơn")
            rows.append((group, label, now, before, change, verdict, fmt))
        table = pd.DataFrame(rows, columns=["Nhóm KPI", "Chỉ số", f"Tháng {month}", "Tháng trước", "Thay đổi", "Đánh giá", "_fmt"])
        row = _h2(ws, row, "Bảng KPI")
        head, last, letters = _table(ws, row, table.drop(columns="_fmt"))
        fmt_map = {"pct": PCT, "int": INT, "money": MONEY, "float": "0.00"}
        for r, fmt in zip(range(head + 1, last + 1), table["_fmt"]):
            for c in ["C", "D", "E"]:
                ws[f"{c}{r}"].number_format = fmt_map[fmt]
        rng = f"F{head + 1}:F{last}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Tốt hơn"'], fill=FILL_GREEN))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Xấu hơn"'], fill=FILL_RED))
        row = last + 2

    trend = kpi_table[["success_tickets", "payment_success_rate", "app_completion_rate"]].tail(12)
    trend.columns = ["Vé thành công", "Tỷ lệ thanh toán thành công", "Tỷ lệ hoàn tất trên app"]
    row = _h2(ws, row, "Xu hướng 12 tháng gần nhất")
    head, last, _ = _table(ws, row, trend, {"Vé thành công": INT, "Tỷ lệ thanh toán thành công": PCT,
                                            "Tỷ lệ hoàn tất trên app": PCT}, index=True)
    line = LineChart()
    line.title, line.height, line.width = "Tỷ lệ thanh toán thành công theo tháng", 7, 18
    line.add_data(Reference(ws, min_col=3, max_col=4, min_row=head, max_row=last), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=head + 1, max_row=last))
    line.y_axis.number_format = "0%"
    ws.add_chart(line, f"H{head}")
    _widths(ws, {"A": 34, "B": 40, "C": 16, "D": 16, "E": 14, "F": 12})


def _marketing_sheet(wb, facts, brief, tables):
    ws = wb.create_sheet(SHEETS["marketing"])
    row = _title(ws, "Marketing · phân nhóm khách & hiệu quả campaign", f"Tuần {facts['week_start']}")
    row = _brief(ws, row, *brief)
    targets = tables["marketing_targets"]
    seg = targets.groupby("segment").agg(**{"Số khách": ("score", "size"), "Điểm quay lại TB": ("score", "mean")})
    seg["Hướng xử lý"] = [facts["marketing"]["segment_actions"].get(s, "") for s in seg.index]
    row = _h2(ws, row, "Phân nhóm khách (luật theo slide báo cáo)")
    _, last, _ = _table(ws, row, seg.sort_values("Số khách", ascending=False).rename_axis("Nhóm"),
                        {"Số khách": INT, "Điểm quay lại TB": "0.000"}, index=True)
    row = last + 2

    camp = tables["campaign_effectiveness"][["new_customers", "returned_90d", "one_time_rate", "discount_per_returned"]]
    camp.columns = ["Khách mới", "Mua lần 2 trong 90 ngày", "Chỉ mua 1 lần", "Giảm giá / 1 khách quay lại"]
    row = _h2(ws, row, "Campaign nào chỉ hút khách một lần? (theo loại)")
    head, last, letters = _table(ws, row, camp.rename_axis("Loại campaign"),
                                 {"Khách mới": INT, "Mua lần 2 trong 90 ngày": PCT, "Chỉ mua 1 lần": PCT,
                                  "Giảm giá / 1 khách quay lại": MONEY}, index=True)
    col = letters["Chỉ mua 1 lần"]
    ws.conditional_formatting.add(f"{col}{head + 1}:{col}{last}", ColorScaleRule(
        start_type="min", start_color="FFFFFF", end_type="max", end_color="F28B82"))
    bar = BarChart()
    bar.type, bar.title, bar.height, bar.width = "bar", "Giảm giá cho mỗi khách quay lại", 6, 14
    gcol = list(letters).index("Giảm giá / 1 khách quay lại") + 1
    bar.add_data(Reference(ws, min_col=gcol, min_row=head, max_row=last), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=head + 1, max_row=last))
    bar.legend = None
    ws.add_chart(bar, f"H{head}")
    row = max(last + 2, head + 14)

    ids = tables["campaign_ids"][["new_customers", "returned_90d", "one_time_rate", "discount_per_returned"]].head(15)
    ids.columns = camp.columns
    row = _h2(ws, row, "Từng campaign (≥ 300 khách mới) - tỷ lệ chỉ mua 1 lần cao nhất")
    head, last, letters = _table(ws, row, ids.rename_axis("campaign_id"),
                                 {"Khách mới": INT, "Mua lần 2 trong 90 ngày": PCT, "Chỉ mua 1 lần": PCT,
                                  "Giảm giá / 1 khách quay lại": MONEY}, index=True)
    col = letters["Chỉ mua 1 lần"]
    ws.conditional_formatting.add(f"{col}{head + 1}:{col}{last}",
                                  CellIsRule(operator="greaterThanOrEqual", formula=["0.95"], fill=FILL_RED))
    row = last + 2

    top = targets.head(500).rename_axis("customer_id").reset_index()
    top = top[["customer_id", "segment", "score", "action", "recency_days", "frequency", "monetary", "promo_share"]]
    top.columns = ["customer_id", "Nhóm", "Điểm quay lại", "Hướng xử lý", "Số ngày từ lần mua gần nhất",
                   "Số lần mua", "Tổng chi tiêu", "Tỷ lệ dùng KM"]
    row = _h2(ws, row, f"Top 500 khách theo điểm quay lại ({facts['marketing']['score_source']}) - đầy đủ trong marketing_targets.csv")
    head, last, _ = _table(ws, row, top, {"Điểm quay lại": "0.000", "Tổng chi tiêu": MONEY, "Tỷ lệ dùng KM": PCT})
    ws.auto_filter.ref = f"A{head}:H{last}"
    _widths(ws, {"A": 26, "B": 24, "C": 16, "D": 60, "E": 20, "F": 12, "G": 14, "H": 14})


def _customer_care_sheet(wb, facts, brief, recovery):
    ws = wb.create_sheet(SHEETS["customer_care"])
    row = _title(ws, "CSKH · danh sách cứu đơn", f"Tuần {facts['week_start']} → {facts['week_end']}")
    ext = facts["error_groups"].get("external", {})
    if ext.get("alert"):
        cell = ws.cell(row=row, column=1, value=f"⚠ Đang có cảnh báo lỗi phía ngân hàng: {ext['rate']} (mức nền {ext['baseline_rate']}). "
                                                "Chủ động hướng dẫn khách đổi sang Ví trong app.")
        cell.fill, cell.font = FILL_RED, Font(bold=True, color="842029")
        row += 2
    row = _brief(ws, row, *brief)
    counts = pd.Series(facts["customer_care"]["errors_by_code"], name="Số lỗi").rename_axis("Mã lỗi").to_frame()
    row = _h2(ws, row, "Số lỗi theo mã")
    _, last, _ = _table(ws, row, counts, index=True)
    row = last + 2
    data = recovery.copy()
    data["time"] = data["time"].dt.strftime("%Y-%m-%d %H:%M")
    data.columns = ["ticket_id", "customer_id", "Thời điểm lỗi", "Phương thức", "status_id", "Lỗi", "Tin nhắn gửi khách"]
    row = _h2(ws, row, f"Danh sách cứu đơn: {len(data):,} khách (lỗi và chưa mua lại tới cuối tuần)")
    head, last, letters = _table(ws, row, data)
    ws.auto_filter.ref = f"A{head}:G{last}"
    ws.freeze_panes = f"A{head + 1}"
    # Lỗi phía ngân hàng (-3, -5): khách không tự sửa được -> tô đỏ để ưu tiên
    ws.conditional_formatting.add(f"A{head + 1}:G{last}",
                                  FormulaRule(formula=[f"OR(${letters['status_id']}{head + 1}=-3,${letters['status_id']}{head + 1}=-5)"],
                                              fill=FILL_RED))
    _widths(ws, {"A": 36, "B": 14, "C": 18, "D": 14, "E": 10, "F": 40, "G": 90})


def _product_sheet(wb, facts, brief, tables, alerts):
    ws = wb.create_sheet(SHEETS["product_it"])
    row = _title(ws, "Product / IT · lỗi thanh toán", f"Tuần {facts['week_start']} so với 8 tuần trước")
    row = _brief(ws, row, *brief)
    week = pd.Timestamp(facts["week_start"])
    ext = alerts[(alerts["error_group"] == "external") & (alerts["week"] <= week)].tail(26)
    trend = ext[["week", "n_tickets", "rate", "baseline", "z", "is_alert"]].copy()
    trend["week"] = trend["week"].dt.date
    trend.columns = ["Tuần", "Lượt thanh toán", "Tỷ lệ lỗi ngân hàng", "Mức nền", "z", "Cảnh báo"]
    row = _h2(ws, row, "Lỗi phía ngân hàng (external) - 26 tuần gần nhất")
    head, last, _ = _table(ws, row, trend, {"Lượt thanh toán": INT, "Tỷ lệ lỗi ngân hàng": PCT, "Mức nền": PCT, "z": "0.0"})
    ws.conditional_formatting.add(f"A{head + 1}:F{last}", FormulaRule(formula=[f"$F{head + 1}=TRUE"], fill=FILL_RED))
    line = LineChart()
    line.title, line.height, line.width = "Tỷ lệ lỗi ngân hàng vs mức nền", 8, 20
    line.add_data(Reference(ws, min_col=3, max_col=4, min_row=head, max_row=last), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=head + 1, max_row=last))
    line.y_axis.number_format = "0%"
    ws.add_chart(line, f"H{head}")
    row = last + 2

    br = tables["product_breakdown"].head(20)[["description", "platform", "paying_method", "errors", "attempts",
                                              "rate", "baseline_rate", "change_pp", "excess_errors"]]
    br.columns = ["Mã lỗi", "Nền tảng", "Phương thức", "Số lỗi", "Lượt thanh toán", "Tỷ lệ tuần này",
                  "Tỷ lệ 8 tuần trước", "Thay đổi (điểm %)", "Lỗi vượt mức"]
    row = _h2(ws, row, "Khoanh vùng: mã lỗi × nền tảng × phương thức")
    head, last, letters = _table(ws, row, br, {"Tỷ lệ tuần này": PCT, "Tỷ lệ 8 tuần trước": PCT,
                                               "Thay đổi (điểm %)": "0.0", "Lỗi vượt mức": "0"})
    col = letters["Thay đổi (điểm %)"]
    ws.conditional_formatting.add(f"{col}{head + 1}:{col}{last}",
                                  CellIsRule(operator="greaterThan", formula=["5"], fill=FILL_RED))
    _widths(ws, {"A": 40, "B": 16, "C": 18, "D": 12, "E": 16, "F": 16, "G": 18, "H": 18, "I": 14})


def _finance_sheet(wb, facts, brief, tables):
    ws = wb.create_sheet(SHEETS["finance"])
    row = _title(ws, "Tài chính · thất thoát do lỗi thanh toán & khuyến mãi", "Đơn vị tiền theo dataset")
    row = _brief(ws, row, *brief)
    losses = tables["monthly_losses"].dropna(how="all").tail(14)
    table = losses[["failed_value", "unrecovered_failed_value", "unrecovered_share",
                    "discount_total", "discount_one_time", "one_time_discount_share"]]
    table.columns = ["Giá trị đơn lỗi", "Không được mua lại (7 ngày)", "% không mua lại",
                     "Tổng giảm giá", "Giảm giá cho khách không quay lại (90 ngày)", "% giảm giá lãng phí"]
    row = _h2(ws, row, "Theo tháng (ô trống = chưa đủ thời gian quan sát)")
    head, last, _ = _table(ws, row, table.rename_axis("Tháng"),
                           {c: (PCT if c.startswith("%") else MONEY) for c in table.columns}, index=True)
    bar = BarChart()
    bar.type, bar.grouping, bar.title, bar.height, bar.width = "col", "clustered", "Thất thoát theo tháng", 8, 22
    bar.add_data(Reference(ws, min_col=3, min_row=head, max_row=last), titles_from_data=True)
    bar.add_data(Reference(ws, min_col=6, min_row=head, max_row=last), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=head + 1, max_row=last))
    ws.add_chart(bar, f"A{last + 2}")
    _widths(ws, {"A": 12, "B": 18, "C": 26, "D": 16, "E": 16, "F": 36, "G": 20})


def _quality_sheet(wb, dq_log):
    ws = wb.create_sheet("Chất lượng dữ liệu")
    row = _title(ws, "Chất lượng dữ liệu đầu vào", "Kết quả kiểm tra từng file ngày (src/quality.py)")
    if dq_log is None or dq_log.empty:
        ws.cell(row=row, column=1, value="Chưa có log - pipeline đang đọc file lịch sử, chưa dùng luồng nạp theo ngày (src/ingest.py).")
        return
    log = dq_log.tail(30).copy()
    head, last, letters = _table(ws, row, log)
    col = letters["status"]
    ws.conditional_formatting.add(f"A{head + 1}:{get_column_letter(len(log.columns))}{last}",
                                  FormulaRule(formula=[f'${col}{head + 1}="failed"'], fill=FILL_RED))
    _widths(ws, {"A": 12, "B": 32, "C": 8, "D": 10, "E": 40, "F": 30, "G": 20})


def build(path, facts, briefs, tables, recovery, alerts, dq_log=None, departments=None):
    """briefs: {dept: (brief, source)}. Ghi file .xlsx tại `path` (đường dẫn hoặc BytesIO).

    departments: chỉ đưa vào sheet của các phòng này (None = tất cả) - dùng để gửi mỗi phòng file riêng.
    """
    departments = list(SHEETS) if departments is None else departments
    wb = Workbook()
    dq_status = "Chưa dùng luồng nạp theo ngày"
    if dq_log is not None and not dq_log.empty:
        last = dq_log.iloc[-1]
        dq_status = f"LỖI - {last['file']}: {last['errors']}" if last["status"] == "failed" else f"Đạt - file cuối {last['file']}"
    _summary_sheet(wb, facts, briefs, recovery, dq_status, departments)
    sheet_builders = {
        "leadership": lambda: _leadership_sheet(wb, facts, briefs["leadership"], tables),
        "marketing": lambda: _marketing_sheet(wb, facts, briefs["marketing"], tables),
        "customer_care": lambda: _customer_care_sheet(wb, facts, briefs["customer_care"], recovery),
        "product_it": lambda: _product_sheet(wb, facts, briefs["product_it"], tables, alerts),
        "finance": lambda: _finance_sheet(wb, facts, briefs["finance"], tables),
    }
    for dept in SHEETS:
        if dept in departments:
            sheet_builders[dept]()
    _quality_sheet(wb, dq_log)
    wb.save(path)
    return path
